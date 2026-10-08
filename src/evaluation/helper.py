import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def parse_prototext(file_path):
    results = []
    current_name = None
    current_metrics = {}

    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Split by evaluation blocks
    blocks = re.split(r'# Evaluating predictions from:', content)[1:]  # skip anything before first #
    for block in blocks:
        # Extract Name
        name_match = re.match(r'\s*(.*?), truth from:', block)
        if name_match:
            current_name = name_match.group(1).strip()
        else:
            current_name = "Unknown"

        # Extract all measures in this block
        measures = re.findall(r'key:\s*"([^"]+)"\s*value:\s*"([^"]+)"', block)
        current_metrics = {"Name": current_name}
        for k, v in measures:
            try:
                v = float(v)
            except ValueError:
                pass
            current_metrics[k] = v

        results.append(current_metrics)

    return results


def save_to_excel(results, output_file="results.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    # Headers dynamically
    all_keys = set(k for r in results for k in r.keys())
    headers = ["Name"] + sorted(k for k in all_keys if k != "Name")
    ws.append(headers)

    # Format header
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Fill rows
    for r in results:
        row = [r.get(h, "") for h in headers]
        ws.append(row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    wb.save(output_file)
    print(f"✅ Results saved to {output_file}")


if __name__ == "__main__":
    input_file = "evaluation.prototext"
    results = parse_prototext(input_file)
    save_to_excel(results)
